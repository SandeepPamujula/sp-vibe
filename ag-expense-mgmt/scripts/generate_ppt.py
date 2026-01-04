
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Design Constants
THEME_COLOR_PRIMARY = RGBColor(0, 51, 102) # Dark Blue
THEME_COLOR_SECONDARY = RGBColor(0, 102, 204) # Lighter Blue
TEXT_COLOR_MAIN = RGBColor(64, 64, 64) # Dark Gray
FONT_TITLE = 'Arial'
FONT_BODY = 'Corbel' # Cleaner sans-serif

def parse_markdown(file_path):
    with open(file_path, 'r') as f:
        content = f.read()

    slides = []
    lines = content.split('\n')
    current_slide = {'title': '', 'content': []}
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('# '):
            if current_slide['title']:
                slides.append(current_slide)
            current_slide = {'title': line[2:], 'content': [], 'type': 'title'}
            
        elif line.startswith('## '):
            if current_slide['title'] or current_slide['content']:
                slides.append(current_slide)
            current_slide = {'title': line[3:], 'content': [], 'type': 'content'}
            
        elif line.startswith('### '):
             current_slide['content'].append({'text': line[4:], 'level': 0, 'bold': True, 'is_header': True})

        elif line.startswith('* ') or line.startswith('- '):
            indent = 0
            if line.startswith('    '):
                indent = 1
            cleaned = re.sub(r'^(\s*)[*-]\s+', '', line)
            current_slide['content'].append({'text': cleaned, 'level': indent})
            
        elif line == '---':
            continue
            
        else:
            if line.startswith('**') and line.endswith('**'):
                 current_slide['content'].append({'text': line.replace('**', ''), 'level': 0, 'bold': True})
            else:
                 current_slide['content'].append({'text': line, 'level': 0})

    if current_slide['title'] or current_slide['content']:
        slides.append(current_slide)
        
    return slides

def style_text_frame(text_frame, content_items):
    text_frame.clear()
    
    for item in content_items:
        p = text_frame.add_paragraph()
        p.text = item['text']
        p.level = item.get('level', 0)
        
        # Font Formatting
        run = p.runs[0] if p.runs else p.add_run()
        run.font.name = FONT_BODY
        run.font.size = Pt(20) if item.get('level', 0) == 0 else Pt(18)
        run.font.color.rgb = TEXT_COLOR_MAIN
        
        if item.get('bold') or item.get('is_header'):
            run.font.bold = True
            if item.get('is_header'):
                run.font.size = Pt(22)
                run.font.color.rgb = THEME_COLOR_SECONDARY
                p.space_before = Pt(12)

        p.space_after = Pt(10)

def create_ppt(slides, output_file):
    prs = Presentation()
    
    # 1. Title Slide
    if slides and slides[0]['type'] == 'title':
        title_slide_data = slides.pop(0)
        slide_layout = prs.slide_layouts[0] 
        slide = prs.slides.add_slide(slide_layout)
        
        # Background Design (Simple Bar)
        left = 0
        top = Inches(6.5)
        width = prs.slide_width
        height = Inches(1)
        shape = slide.shapes.add_shape(1, left, top, width, height) # 1 = msoShapeRectangle
        shape.fill.solid()
        shape.fill.fore_color.rgb = THEME_COLOR_PRIMARY
        shape.line.fill.background()

        title = slide.shapes.title
        title.text = title_slide_data['title']
        title.text_frame.paragraphs[0].font.name = FONT_TITLE
        title.text_frame.paragraphs[0].font.size = Pt(44)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = THEME_COLOR_PRIMARY
        
        subtitle = slide.placeholders[1]
        subtitle.text = "\n".join([item['text'] for item in title_slide_data['content']])
        for paragraph in subtitle.text_frame.paragraphs:
            paragraph.font.name = FONT_BODY
            paragraph.font.size = Pt(24)
            paragraph.font.color.rgb = TEXT_COLOR_MAIN

    # 2. Content Slides
    for slide_data in slides:
        slide_layout = prs.slide_layouts[1] # Title and Content
        slide = prs.slides.add_slide(slide_layout)
        
        # Customize Title
        title = slide.shapes.title
        title.text = slide_data['title']
        # Move Title up and align left
        title.top = Inches(0.5)
        title.left = Inches(0.5)
        title.width = Inches(9)
        
        TitlePara = title.text_frame.paragraphs[0]
        TitlePara.font.name = FONT_TITLE
        TitlePara.font.size = Pt(32)
        TitlePara.font.bold = True
        TitlePara.font.color.rgb = THEME_COLOR_PRIMARY
        TitlePara.alignment = PP_ALIGN.LEFT
        
        # Add a decorative line
        line = slide.shapes.add_shape(
            1, Inches(0.5), Inches(1.2), Inches(9), Inches(0.05)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = THEME_COLOR_SECONDARY
        line.line.fill.background()

        # Customize Body
        body = slide.placeholders[1]
        body.top = Inches(1.5)
        body.left = Inches(0.5)
        body.width = Inches(9)
        body.height = Inches(5.5)
        
        style_text_frame(body.text_frame, slide_data['content'])

    prs.save(output_file)
    print(f"Presentation saved to {output_file}")

if __name__ == "__main__":
    input_md = "docs/presentation_architects.md"
    output_pptx = "docs/presentation_architects_styled.pptx"
    
    if not os.path.exists(input_md):
        print(f"Error: {input_md} not found.")
    else:
        slides = parse_markdown(input_md)
        create_ppt(slides, output_pptx)
